# -*- coding: utf-8 -*-
import io
import re
import logging
from datetime import datetime

from odoo import fields, models, _
from odoo.exceptions import UserError

_logger = logging.getLogger(__name__)

try:
    import openpyxl
except ImportError:
    _logger.warning('openpyxl library not found. Install it with: pip install openpyxl')
    openpyxl = None


class CodReconciliationViettelpost(models.Model):
    _inherit = 'cod.reconciliation'



    def _parse_reconciliation_file(self, file_data, filename):
        """Parse ViettelPost COD reconciliation Excel file (BangKeChiCod).

        Overrides the base method to implement ViettelPost-specific parsing.
        """
        if not openpyxl:
            raise UserError(_(
                'The openpyxl library is required to import Excel files. '
                'Please install it with: pip install openpyxl'
            ))

        try:
            wb = openpyxl.load_workbook(io.BytesIO(file_data), data_only=True)
            ws = wb.active
        except Exception as e:
            raise UserError(_('Failed to read Excel file: %s', str(e)))

        # Find section rows
        sections = self._vtp_find_section_rows(ws)
        _logger.info('ViettelPost - Detected Excel sections: %s', sections)

        if not sections.get('cod_data_start'):
            raise UserError(_(
                'Could not detect the COD data section in the Excel file. '
                'Please make sure the file follows the ViettelPost reconciliation format '
                '(BangKeChiCod).'
            ))

        # Parse header info
        header_info = self._vtp_parse_header_info(ws)

        # Parse COD section
        cod_lines = self._vtp_parse_cod_section(ws, sections)
        _logger.info('ViettelPost - Parsed %d COD lines', len(cod_lines))

        if not cod_lines:
            raise UserError(_('No COD data lines found in the Excel file.'))

        # Parse fee section
        fee_data = self._vtp_parse_fee_section(ws, sections)
        _logger.info('ViettelPost - Parsed %d fee entries', len(fee_data))

        # Parse summary
        summary = self._vtp_parse_summary(ws, sections)
        _logger.info('ViettelPost - Parsed summary: %s', summary)

        # Merge COD lines with fee data
        result_lines = []
        for line in cod_lines:
            bill_code = line['carrier_bill_code']

            # Merge fee data if available
            if bill_code in fee_data:
                line.update(fee_data[bill_code])
                if fee_data[bill_code].get('carrier_note') and not line.get('carrier_note'):
                    line['carrier_note'] = fee_data[bill_code]['carrier_note']

            result_lines.append(line)

        # Also add fee-only lines (bills present in fee section but not in COD section)
        cod_bill_codes = {l['carrier_bill_code'] for l in cod_lines}
        for bill_code, fee_info in fee_data.items():
            if bill_code not in cod_bill_codes:
                result_lines.append({
                    'carrier_bill_code': bill_code,
                    'carrier_cod_amount': 0.0,
                    **fee_info,
                })

        return {
            'lines': result_lines,
            'carrier_total_cod': summary.get('carrier_total_cod', 0),
            'carrier_total_shipping_fee': summary.get('carrier_total_shipping_fee', 0),
            'carrier_net_amount': summary.get('carrier_net_amount', 0),
            'period_date': header_info.get('period_date', False),

        }

    # ------------------------------------------------------------------
    # ViettelPost-specific parsing helpers
    # ------------------------------------------------------------------

    @staticmethod
    def _vtp_parse_date(value):
        """Parse date from various formats used by ViettelPost."""
        if not value:
            return False
        if isinstance(value, datetime):
            return value.date()
        if isinstance(value, str):
            for fmt in ('%d/%m/%Y', '%Y-%m-%d', '%d-%m-%Y', '%d.%m.%Y'):
                try:
                    return datetime.strptime(value.strip(), fmt).date()
                except (ValueError, AttributeError):
                    continue
        return False

    @staticmethod
    def _vtp_parse_float(value):
        """Parse float value, handling None and string formats."""
        if not value:
            return 0.0
        if isinstance(value, (int, float)):
            return float(value)
        if isinstance(value, str):
            try:
                cleaned = value.replace(',', '').replace(' ', '').strip()
                return float(cleaned)
            except (ValueError, AttributeError):
                return 0.0
        return 0.0

    def _vtp_find_section_rows(self, ws):
        """Find the start rows of the two data sections and the summary section."""
        result = {
            'cod_header_row': None,
            'cod_data_start': None,
            'fee_header_row': None,
            'fee_data_start': None,
            'summary_start': None,
        }

        for row in ws.iter_rows(min_row=1, max_row=ws.max_row, max_col=10, values_only=False):
            row_num = row[0].row
            cell_a = str(row[0].value or '').strip()

            # Look for section headers containing "CHI TIẾT"
            for cell in row:
                cell_val = str(cell.value or '').strip().upper()
                if 'CHI TIẾT SỐ TIỀN COD' in cell_val:
                    pass  # Just a marker, we look for STT row after this
                elif 'CHI TIẾT TIỀN CƯỚC' in cell_val:
                    pass  # Just a marker
                elif 'KẾT LUẬN ĐỐI SOÁT' in cell_val:
                    result['summary_start'] = row_num

            # Look for header rows with "STT"
            if cell_a.upper() == 'STT':
                if not result['cod_header_row']:
                    result['cod_header_row'] = row_num
                    result['cod_data_start'] = row_num + 1
                elif not result['fee_header_row']:
                    result['fee_header_row'] = row_num
                    result['fee_data_start'] = row_num + 1

        return result

    @staticmethod
    def _vtp_is_total_row(row_values):
        """Check if a row is a TỔNG CỘNG (total) row."""
        for val in row_values:
            if val and 'TỔNG CỘNG' in str(val).upper():
                return True
        return False

    @staticmethod
    def _vtp_is_data_row(row_values):
        """Check if this is a valid data row (has a numeric STT or valid bill code)."""
        stt = row_values[0]
        bill_code = row_values[1]

        if stt is not None:
            try:
                int(float(str(stt)))
                return True
            except (ValueError, TypeError):
                pass

        if bill_code and str(bill_code).strip():
            try:
                int(float(str(bill_code).strip()))
                return True
            except (ValueError, TypeError):
                pass

        return False

    def _vtp_parse_cod_section(self, ws, sections):
        """Parse ViettelPost COD amount section.

        Columns: A=STT, B=Số BILL, C=Ngày gửi, D=Dịch vụ, E=Ngày phát thành công, F=Số tiền COD, G=Ghi chú
        """
        lines = []
        if not sections.get('cod_data_start'):
            return lines

        start_row = sections['cod_data_start']
        end_row = sections.get('fee_header_row') or ws.max_row

        for row in ws.iter_rows(min_row=start_row, max_row=end_row, values_only=False):
            row_values = [cell.value for cell in row]

            if self._vtp_is_total_row(row_values):
                break

            if not self._vtp_is_data_row(row_values):
                continue

            bill_code = str(row_values[1]).strip() if row_values[1] else False
            if not bill_code:
                continue

            line = {
                'sequence': int(float(str(row_values[0]))) if row_values[0] else 0,
                'carrier_bill_code': bill_code,
                'carrier_send_date': self._vtp_parse_date(row_values[2]),
                'carrier_service': str(row_values[3]).strip() if row_values[3] else '',
                'carrier_delivery_date': self._vtp_parse_date(row_values[4]),
                'carrier_cod_amount': self._vtp_parse_float(row_values[5]),
                'carrier_note': str(row_values[6]).strip() if row_values[6] else '',
            }
            lines.append(line)

        return lines

    def _vtp_parse_fee_section(self, ws, sections):
        """Parse ViettelPost shipping fee section.

        Columns: A=STT, B=Số BILL, C=Ngày gửi, D=Dịch vụ, E=Trọng lượng,
                 F=Cước phí, G=Cước đã thu, H=Giảm giá, I=Tổng số tiền, J=Ghi chú
        """
        fees = {}
        if not sections.get('fee_data_start'):
            return fees

        start_row = sections['fee_data_start']
        end_row = sections.get('summary_start') or ws.max_row

        for row in ws.iter_rows(min_row=start_row, max_row=end_row, values_only=False):
            row_values = [cell.value for cell in row]

            if self._vtp_is_total_row(row_values):
                break

            if not self._vtp_is_data_row(row_values):
                continue

            bill_code = str(row_values[1]).strip() if row_values[1] else False
            if not bill_code:
                continue

            fees[bill_code] = {
                'carrier_weight': self._vtp_parse_float(row_values[4]),
                'carrier_shipping_fee': self._vtp_parse_float(row_values[5]),
                'carrier_collected_fee': self._vtp_parse_float(row_values[6]),
                'carrier_discount': self._vtp_parse_float(row_values[7]),
                'carrier_total_fee': self._vtp_parse_float(row_values[8]),
            }
            if len(row_values) > 9 and row_values[9]:
                fees[bill_code]['carrier_note'] = str(row_values[9]).strip()

        return fees

    def _vtp_parse_summary(self, ws, sections):
        """Parse the summary section at the bottom of ViettelPost report."""
        summary = {
            'carrier_total_cod': 0.0,
            'carrier_total_shipping_fee': 0.0,
            'carrier_net_amount': 0.0,
        }

        if not sections.get('summary_start'):
            return summary

        for row in ws.iter_rows(
            min_row=sections['summary_start'],
            max_row=min(sections['summary_start'] + 10, ws.max_row),
            values_only=False,
        ):
            row_values = [cell.value for cell in row]
            row_text = ' '.join(str(v or '') for v in row_values).upper()

            if 'SỐ TIỀN COD PHẢI TRẢ' in row_text or '1. SỐ TIỀN COD' in row_text:
                for val in row_values:
                    amount = self._vtp_parse_float(val)
                    if amount > 0:
                        summary['carrier_total_cod'] = amount
                        break
            elif 'SỐ TIỀN CƯỚC CPN' in row_text or '2. SỐ TIỀN CƯỚC' in row_text:
                for val in row_values:
                    amount = self._vtp_parse_float(val)
                    if amount > 0:
                        summary['carrier_total_shipping_fee'] = amount
                        break
            elif 'SỐ TIỀN CÒN LẠI' in row_text or '3. SỐ TIỀN CÒN LẠI' in row_text:
                for val in row_values:
                    amount = self._vtp_parse_float(val)
                    if amount > 0:
                        summary['carrier_net_amount'] = amount
                        break

        return summary

    def _vtp_parse_header_info(self, ws):
        """Parse ViettelPost header info (customer code, name, tax code, period)."""
        info = {}
        for row in ws.iter_rows(min_row=1, max_row=12, values_only=False):
            row_values = [cell.value for cell in row]
            row_text = ' '.join(str(v or '') for v in row_values)

            for i, val in enumerate(row_values):
                val_str = str(val or '')

                if 'Mã khách hàng' in val_str or 'MÃ KHÁCH HÀNG' in val_str.upper():
                    if ':' in val_str:
                        info['vtp_customer_code'] = val_str.split(':', 1)[1].strip()
                    elif i + 1 < len(row_values) and row_values[i + 1]:
                        info['vtp_customer_code'] = str(row_values[i + 1]).strip()

                if 'Tên khách hàng' in val_str or 'TÊN KHÁCH HÀNG' in val_str.upper():
                    if ':' in val_str:
                        info['vtp_customer_name'] = val_str.split(':', 1)[1].strip()
                    elif i + 1 < len(row_values) and row_values[i + 1]:
                        info['vtp_customer_name'] = str(row_values[i + 1]).strip()

                if 'Mã số thuế' in val_str or 'MÃ SỐ THUẾ' in val_str.upper():
                    if ':' in val_str:
                        info['vtp_tax_code'] = val_str.split(':', 1)[1].strip()
                    elif i + 1 < len(row_values) and row_values[i + 1]:
                        info['vtp_tax_code'] = str(row_values[i + 1]).strip()

            # Look for period date in format "Ngày XX tháng XX năm XXXX"
            for val in row_values:
                val_str = str(val or '')
                if 'tháng' in val_str and 'năm' in val_str:
                    match = re.search(r'(\d{1,2})\s*tháng\s*(\d{1,2})\s*năm\s*(\d{4})', val_str)
                    if match:
                        day, month, year = match.groups()
                        try:
                            info['period_date'] = datetime(int(year), int(month), int(day)).date()
                        except ValueError:
                            pass

        return info


